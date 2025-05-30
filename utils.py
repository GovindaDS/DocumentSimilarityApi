# utils.py

import numpy as np
import os
from pdfminer.high_level import extract_text
from docx import Document
import openpyxl

def extract_text_from_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text(file_path)
    elif ext == ".docx":
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs])
    elif ext in [".xlsx", ".xls"]:
        wb = openpyxl.load_workbook(file_path)
        text = ""
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) if cell else "" for cell in row]) + "\n"
        return text
    else:
        raise ValueError("Unsupported file type")

def split_document_into_chunks(document_text, max_tokens=512, overlap=50, tokenizer=None):
    if tokenizer is None:
        words = document_text.split()
        chunk_size_words = int(max_tokens * 0.75)
        overlap_words = int(overlap * 0.75)
        chunks = []
        i = 0
        while i < len(words):
            chunk = ' '.join(words[i: i + chunk_size_words])
            chunks.append(chunk)
            i += chunk_size_words - overlap_words
            if i < 0:
                i = 0
        return chunks

    tokens = tokenizer.encode(document_text, add_special_tokens=False)
    chunks = []
    i = 0
    while i < len(tokens):
        chunk_tokens = tokens[i: i + max_tokens]
        chunks.append(tokenizer.decode(chunk_tokens))
        i += max_tokens - overlap
        if i < 0:
            i = 0
    return chunks

def get_document_embedding(document_text, model):
    tokenizer = model.tokenizer
    max_tokens = model.max_seq_length
    chunks = split_document_into_chunks(document_text, max_tokens=max_tokens, overlap=max_tokens // 10, tokenizer=tokenizer)
    if not chunks:
        return np.zeros(model.get_sentence_embedding_dimension())
    chunk_embeddings = model.encode(chunks, convert_to_tensor=True, show_progress_bar=False)
    return np.mean(chunk_embeddings.cpu().numpy(), axis=0)
